from io import BytesIO
from datetime import datetime


def as_markdown(title: str, content: str) -> str:
    return (
        f"# {title}\n\n"
        f"Generated: {datetime.now().isoformat(timespec='minutes')}\n\n"
        f"{content}\n\n"
        "---\nHuman review required. Do not use identifiable patient data."
    )


def as_txt(title: str, content: str) -> bytes:
    return as_markdown(title, content).encode("utf-8")


def as_docx(title: str, content: str) -> bytes:
    try:
        from docx import Document

        doc = Document()
        doc.add_heading(title, 0)
        for block in content.split("\n\n"):
            doc.add_paragraph(block)
        doc.add_heading("Safety note", level=1)
        doc.add_paragraph("Human review required. Do not use identifiable patient information.")
        bio = BytesIO()
        doc.save(bio)
        return bio.getvalue()
    except Exception:
        return as_txt(title, content)


def as_pdf(title: str, content: str) -> bytes:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

        bio = BytesIO()
        doc = SimpleDocTemplate(bio, pagesize=A4)
        styles = getSampleStyleSheet()
        story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
        for block in content.split("\n\n"):
            safe = (
                block.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace("\n", "<br/>")
            )
            story += [Paragraph(safe, styles["BodyText"]), Spacer(1, 8)]
        story += [Paragraph("Human review required. Do not use identifiable patient data.", styles["Italic"])]
        doc.build(story)
        return bio.getvalue()
    except Exception:
        return as_txt(title, content)
